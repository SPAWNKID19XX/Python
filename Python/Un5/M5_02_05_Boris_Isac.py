import openpyxl
from openpyxl.chart import BarChart, Reference

file = openpyxl.load_workbook("02_Excel_data.xlsx")
olimpName = file.sheetnames[2]
olimpSheet = file[olimpName]

graph = BarChart()
graph.title="Medalhas Olimpicas"

countries = Reference(olimpSheet,min_col=1, min_row=2, max_col=1, max_row=7)
selected = Reference(olimpSheet,min_col=5, min_row=2, max_col=5, max_row=7)
graph.x_axis.title = "Paises"
graph.y_axis.title = "Medalhas"
graph.style = 3
graph.legend = None
graph.varyColors = True
#print(countries)#Cheking select cells
#print(selected)#Cheking select cells
graph.add_data(selected, titles_from_data=True)
graph.set_categories(countries)
olimpSheet.add_chart(graph, "G1")
file.save("02_Excel_data.xlsx")
file.close()

